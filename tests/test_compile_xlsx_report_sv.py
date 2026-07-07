import importlib.util
from pathlib import Path

import openpyxl

SCRIPT = Path(__file__).parent.parent / "workflow" / "scripts" / "compile_xlsx_report_sv.py"

# Single-sample VCF (HG002 — one sample column)
VCF_SINGLE = """\
##fileformat=VCFv4.2
##INFO=<ID=SVTYPE,Number=1,Type=String,Description="SV type">
##INFO=<ID=SVLEN,Number=1,Type=Integer,Description="SV length">
##INFO=<ID=END,Number=1,Type=Integer,Description="End position">
##INFO=<ID=VAF,Number=1,Type=Float,Description="Variant allele frequency">
##INFO=<ID=COVERAGE,Number=.,Type=Integer,Description="Coverage">
##FORMAT=<ID=GT,Number=1,Type=String,Description="Genotype">
##FORMAT=<ID=DR,Number=1,Type=Integer,Description="Ref reads">
##FORMAT=<ID=DV,Number=1,Type=Integer,Description="Alt reads">
#CHROM\tPOS\tID\tREF\tALT\tQUAL\tFILTER\tINFO\tFORMAT\tHG002
chr1\t999\tSniffles2.DEL.1\tN\t<DEL>\t60\tPASS\tSVTYPE=DEL;SVLEN=-500;END=1500;VAF=0.5;COVERAGE=30,28\tGT:DR:DV\t0/1:15:15
chr2\t1999\tSniffles2.INS.1\tN\t<INS>\t60\tPASS\tSVTYPE=INS;SVLEN=300;END=2000;VAF=0.9;COVERAGE=20\tGT:DR:DV\t1/1:2:18
chr3\t2999\tSniffles2.XYZ.1\tN\t<XYZ>\t60\tPASS\tSVTYPE=XYZ;END=3100\tGT:DR:DV\t0/1:10:5
"""

# Joint VCF (HG002_T tumor + HG002_N normal — two sample columns)
VCF_JOINT = """\
##fileformat=VCFv4.2
##INFO=<ID=SVTYPE,Number=1,Type=String,Description="SV type">
##INFO=<ID=SVLEN,Number=1,Type=Integer,Description="SV length">
##INFO=<ID=END,Number=1,Type=Integer,Description="End position">
##INFO=<ID=VAF,Number=1,Type=Float,Description="Variant allele frequency">
##INFO=<ID=COVERAGE,Number=.,Type=Integer,Description="Coverage">
##FORMAT=<ID=GT,Number=1,Type=String,Description="Genotype">
##FORMAT=<ID=DR,Number=1,Type=Integer,Description="Ref reads">
##FORMAT=<ID=DV,Number=1,Type=Integer,Description="Alt reads">
#CHROM\tPOS\tID\tREF\tALT\tQUAL\tFILTER\tINFO\tFORMAT\tHG002_T\tHG002_N
chr1\t999\tSniffles2.DEL.1\tN\t<DEL>\t60\tPASS\tSVTYPE=DEL;SVLEN=-500;END=1500;VAF=0.5;COVERAGE=30,28\tGT:DR:DV\t0/1:10:8\t0/0:20:1
chr2\t1999\tSniffles2.INS.1\tN\t<INS>\t60\tPASS\tSVTYPE=INS;SVLEN=300;END=2000;VAF=0.9;COVERAGE=20\tGT:DR:DV\t1/1:1:17\t0/0:18:0
chr3\t2999\tSniffles2.XYZ.1\tN\t<XYZ>\t60\tPASS\tSVTYPE=XYZ;END=3100\tGT:DR:DV\t0/1:8:4\t0/0:15:0
"""


def make_vcf(tmp_path, content, name="sv.vcf"):
    vcf = tmp_path / name
    vcf.write_text(content)
    return str(vcf)


def load_module():
    spec = importlib.util.spec_from_file_location("xlsx_sv", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# ── Single-sample tests (existing behaviour) ──────────────────────────────────

def test_parse_sv_vcf_keeps_allowed_types(tmp_path):
    mod = load_module()
    rows, _ = mod.parse_sv_vcf(make_vcf(tmp_path, VCF_SINGLE))
    assert sorted(r["SVTYPE"] for r in rows) == ["DEL", "INS"]  # XYZ filtered out


def test_parse_sv_vcf_fields(tmp_path):
    mod = load_module()
    rows, _ = mod.parse_sv_vcf(make_vcf(tmp_path, VCF_SINGLE))
    del_row = next(r for r in rows if r["SVTYPE"] == "DEL")
    assert del_row["CHROM"] == "chr1"
    assert del_row["POS"] == 999
    assert del_row["SVLEN"] == -500
    assert del_row["GT"] == "0/1"
    assert del_row["DR"] == 15
    assert del_row["DV"] == 15
    assert abs(del_row["VAF"] - 0.5) < 1e-6
    assert del_row["COVERAGE_T"] == "30,28"


def test_write_xlsx_creates_file(tmp_path):
    mod = load_module()
    rows, sample_order = mod.parse_sv_vcf(make_vcf(tmp_path, VCF_SINGLE))
    out = tmp_path / "out.xlsx"
    mod.write_xlsx(rows, str(out), sample="HG002", software_versions={"sniffles2": "2.2"}, sample_order=sample_order)
    assert out.exists() and out.stat().st_size > 0


def test_write_xlsx_columns(tmp_path):
    mod = load_module()
    rows, sample_order = mod.parse_sv_vcf(make_vcf(tmp_path, VCF_SINGLE))
    out = tmp_path / "out.xlsx"
    mod.write_xlsx(rows, str(out), sample="HG002", software_versions={"sniffles2": "2.2"}, sample_order=sample_order)
    wb = openpyxl.load_workbook(str(out), read_only=True)
    headers = [cell.value for cell in next(wb["SV"].rows)]
    for col in ("CHROM", "POS", "SVTYPE", "GT", "DR", "DV", "VAF", "COVERAGE_T"):
        assert col in headers, f"{col} missing from SV sheet"
    # joint-mode suffixed columns must NOT appear in single-sample output
    assert "GT_T" not in headers
    assert "GT_N" not in headers


# ── Joint / multi-sample tests ─────────────────────────────────────────────────

def test_joint_parse_keeps_allowed_types(tmp_path):
    mod = load_module()
    rows, _ = mod.parse_sv_vcf(
        make_vcf(tmp_path, VCF_JOINT, "joint.vcf"),
        sample_names=["HG002_T", "HG002_N"],
    )
    assert sorted(r["SVTYPE"] for r in rows) == ["DEL", "INS"]


def test_joint_parse_sample_order_matches_header(tmp_path):
    mod = load_module()
    _, sample_order = mod.parse_sv_vcf(
        make_vcf(tmp_path, VCF_JOINT, "joint.vcf"),
        sample_names=["HG002_T", "HG002_N"],
    )
    # VCF_JOINT header lists #CHROM ... HG002_T HG002_N — SUPP_VEC bit order follows this.
    assert sample_order == ["HG002_T", "HG002_N"]


def test_joint_parse_per_sample_columns(tmp_path):
    mod = load_module()
    rows, _ = mod.parse_sv_vcf(
        make_vcf(tmp_path, VCF_JOINT, "joint.vcf"),
        sample_names=["HG002_T", "HG002_N"],
    )
    del_row = next(r for r in rows if r["SVTYPE"] == "DEL")
    # Tumor columns
    assert del_row["GT_T"] == "0/1"
    assert del_row["DR_T"] == 10
    assert del_row["DV_T"] == 8
    # Normal columns
    assert del_row["GT_N"] == "0/0"
    assert del_row["DR_N"] == 20
    assert del_row["DV_N"] == 1
    # Shared INFO fields
    assert del_row["COVERAGE_T"] == "30,28"
    # Joint mode computes VAF as DV/(DR+DV) per sample, not from INFO/VAF (rounded to 4 dp).
    assert abs(del_row["VAF_T"] - round(8 / 18, 4)) < 1e-6
    # Flat (unsuffixed) FORMAT columns must not appear
    assert "GT" not in del_row
    assert "DR" not in del_row


def test_joint_write_xlsx_columns(tmp_path):
    mod = load_module()
    rows, sample_order = mod.parse_sv_vcf(
        make_vcf(tmp_path, VCF_JOINT, "joint.vcf"),
        sample_names=["HG002_T", "HG002_N"],
    )
    out = tmp_path / "joint.xlsx"
    mod.write_xlsx(rows, str(out), sample="HG002", software_versions={"sniffles2": "2.2"}, sample_order=sample_order)
    wb = openpyxl.load_workbook(str(out), read_only=True)
    headers = [cell.value for cell in next(wb["SV"].rows)]
    for col in ("CHROM", "POS", "SVTYPE",
                "GT_T", "DR_T", "DV_T", "VAF_T",
                "GT_N", "DR_N", "DV_N", "VAF_N",
                "COVERAGE_T"):
        assert col in headers, f"{col} missing from joint SV sheet"


def test_joint_write_xlsx_info_sheet_has_supp_vec_order(tmp_path):
    mod = load_module()
    rows, sample_order = mod.parse_sv_vcf(
        make_vcf(tmp_path, VCF_JOINT, "joint.vcf"),
        sample_names=["HG002_T", "HG002_N"],
    )
    out = tmp_path / "joint.xlsx"
    mod.write_xlsx(rows, str(out), sample="HG002", software_versions={"sniffles2": "2.2"}, sample_order=sample_order)
    wb = openpyxl.load_workbook(str(out), read_only=True)
    info_rows = list(wb["info"].rows)
    kv = {row[0].value: row[1].value for row in info_rows[1:]}  # skip header row
    # Only the N/T type suffix is shown, not the full sample name.
    assert kv["SUPP_VEC sample order"] == "1=T,2=N"


def test_joint_write_xlsx_info_sheet_supp_vec_order_follows_header_order(tmp_path):
    mod = load_module()
    # Header lists normal before tumor here — the "1=N,2=T" value must reflect
    # this actual header order, not the sample_names arg order passed in.
    vcf_n_first = VCF_JOINT.replace("HG002_T\tHG002_N", "HG002_N\tHG002_T").replace(
        "0/1:10:8\t0/0:20:1", "0/0:20:1\t0/1:10:8"
    ).replace("1/1:1:17\t0/0:18:0", "0/0:18:0\t1/1:1:17").replace(
        "0/1:8:4\t0/0:15:0", "0/0:15:0\t0/1:8:4"
    )
    rows, sample_order = mod.parse_sv_vcf(
        make_vcf(tmp_path, vcf_n_first, "joint_n_first.vcf"),
        sample_names=["HG002_T", "HG002_N"],
    )
    out = tmp_path / "joint_n_first.xlsx"
    mod.write_xlsx(rows, str(out), sample="HG002", software_versions={"sniffles2": "2.2"}, sample_order=sample_order)
    wb = openpyxl.load_workbook(str(out), read_only=True)
    info_rows = list(wb["info"].rows)
    kv = {row[0].value: row[1].value for row in info_rows[1:]}
    assert kv["SUPP_VEC sample order"] == "1=N,2=T"


def test_supp_vec_order_strips_sample_name_with_underscore_prefix(tmp_path):
    mod = load_module()
    vcf = VCF_JOINT.replace("HG002_T", "PAT_1234_T").replace("HG002_N", "PAT_1234_N")
    rows, sample_order = mod.parse_sv_vcf(
        make_vcf(tmp_path, vcf, "joint_underscored.vcf"),
        sample_names=["PAT_1234_T", "PAT_1234_N"],
    )
    out = tmp_path / "joint_underscored.xlsx"
    mod.write_xlsx(rows, str(out), sample="PAT_1234", software_versions={"sniffles2": "2.2"}, sample_order=sample_order)
    wb = openpyxl.load_workbook(str(out), read_only=True)
    info_rows = list(wb["info"].rows)
    kv = {row[0].value: row[1].value for row in info_rows[1:]}
    assert kv["SUPP_VEC sample order"] == "1=T,2=N"


def test_single_sample_write_xlsx_info_sheet_has_no_supp_vec_order(tmp_path):
    mod = load_module()
    rows, sample_order = mod.parse_sv_vcf(make_vcf(tmp_path, VCF_SINGLE))
    out = tmp_path / "out.xlsx"
    mod.write_xlsx(rows, str(out), sample="HG002", software_versions={"sniffles2": "2.2"}, sample_order=sample_order)
    wb = openpyxl.load_workbook(str(out), read_only=True)
    info_rows = list(wb["info"].rows)
    keys = [row[0].value for row in info_rows[1:]]
    assert "SUPP_VEC sample order" not in keys
